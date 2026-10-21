"""Quality regressions: factual focus, extraction context, boundaries and privacy.

These fixtures are curated expectations, not an open-ended accuracy benchmark.
Hosted-model tests mock its wire response; no API calls or charges are made.
"""
import io
import json
import tempfile
import time
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from app import create_app
from backend.extraction import DocumentText, TextUnit, extract_document
from backend.generation import extractive
from backend import cloud
from backend.errors import UserError
from backend.smart import sentence_parts, passages
from fixtures import docx_bytes


def document(*paragraphs):
    return DocumentText([TextUnit(text, f'Paragraph {i+1}') for i, text in enumerate(paragraphs)], [])


class QualityTests(unittest.TestCase):
    def test_atomic_definitions_and_function_subjects(self):
        cards = extractive(document('A variable is a name that refers to an object in Python.',
                                    'The return statement sends a result from a function to its caller.',
                                    'The len function returns the number of items in an object.'), 20)
        self.assertEqual([(c['question'], c['answer']) for c in cards], [
            ('What is a variable?', 'A name that refers to an object in Python.'),
            ('What does the return statement send?', 'A result from a function to its caller.'),
            ('What does the len function return?', 'The number of items in an object.')])

    def test_textbook_contrasts_causes_and_qualifiers(self):
        cards = extractive(document(
            'Osmosis is the movement of water through a selectively permeable membrane.',
            'Enzymes are proteins that speed up chemical reactions.',
            'The rate of diffusion increases because particles move faster at higher temperatures.',
            'Arteries carry blood away from the heart.',
            'Mitosis is cell division, whereas meiosis is reduction division.'), 20)
        self.assertEqual(len(cards), 5)
        by_kind = {c['kind']: c for c in cards}
        self.assertEqual(by_kind['cause']['question'], 'Why does the rate of diffusion increase?')
        self.assertIn('higher temperatures', by_kind['cause']['answer'])
        self.assertEqual(by_kind['comparison']['question'], 'How do Mitosis and meiosis differ?')
        self.assertIn('whereas', by_kind['comparison']['answer'])
        self.assertFalse(any('notes' in c['question'] or 'passage' in c['question'] for c in cards))

    def test_office_notes_and_short_definitions(self):
        cards = extractive(document('SLA: an agreement defining service quality and response times.',
                                    'The support team requires a ticket number before investigating an incident.',
                                    'API: Application Programming Interface.'), 20)
        self.assertEqual(len(cards), 3)
        self.assertTrue(any(c['answer'] == 'Application Programming Interface.' for c in cards))
        self.assertTrue(any('before investigating' in c['answer'] for c in cards))

    def test_do_not_turn_pronouns_or_furniture_into_cards(self):
        cards = extractive(document('Chapter 1', 'Contents', 'Photosynthesis ........ 12',
                                    'This is an essential process for living organisms.',
                                    'They are proteins that speed up reactions.',
                                    'Copyright 2026 All rights reserved.',
                                    'In this chapter you will learn about the structure of a cell.',
                                    'A nucleus is an organelle that contains genetic material.'), 20)
        self.assertEqual(len(cards), 1)
        self.assertEqual(cards[0]['question'], 'What is a nucleus?')

    def test_cloze_answer_is_only_the_gap_and_keeps_units(self):
        cards = extractive(document('Normal human body temperature is 37 degrees Celsius.',
                                    'A list is an ordered mutable collection of items.'), 20, style='cloze')
        self.assertEqual(len(cards), 2)
        self.assertEqual(cards[0]['answer'], '37')
        for c in cards:
            self.assertEqual(c['question'].count('_____'), 1)
            self.assertNotIn(c['answer'].lower(), c['question'].lower())
            self.assertLess(len(c['answer'].split()), 6)
        self.assertIn('degrees Celsius', cards[0]['question'])

    def test_no_padding_and_near_duplicate_suppression(self):
        cards = extractive(document('A list is an ordered mutable collection of items.',
                                    'A list is an ordered mutable collection of items.',
                                    'A list is an ordered mutable collection of elements.'), 50)
        self.assertEqual(len(cards), 1)

    def test_sentences_preserve_abbreviations_and_decimals(self):
        result = sentence_parts('Dr. Smith measures 3.14 grams. A sample is a representative subset.')
        self.assertEqual(len(result), 2)
        self.assertTrue(result[0].startswith('Dr. Smith'))
        self.assertIn('3.14', result[0])

    def test_docx_heading_and_table_column_semantics(self):
        doc = Document(); doc.add_heading('Algorithms', 1)
        table = doc.add_table(rows=3, cols=3)
        for row, cells in zip(table.rows, [('Method', 'Time complexity', 'Space complexity'),
                                         ('Binary search', 'O(log n)', 'O(1)'),
                                         ('Linear search', 'O(n)', 'O(1)')]):
            for cell, value in zip(row.cells, cells): cell.text = value
        stream=io.BytesIO();doc.save(stream)
        parsed=extract_document(stream.getvalue(), 'algorithms.docx')
        cards=extractive(parsed,20)
        self.assertEqual(len(cards),4)
        self.assertIn('Algorithms', cards[0]['tags'])
        self.assertTrue(any(c['question']=='For Binary search, what is the time complexity?' and c['answer']=='O(log n).' for c in cards))
        self.assertFalse(any('What is Method?' == c['question'] for c in cards))

    def test_explicit_question_answer_and_ordered_steps(self):
        cards=extractive(document('Question: What does HTTP stand for?', 'Answer: Hypertext Transfer Protocol.',
                                  'Steps of the review process:\n1. Read the request\n2. Check the evidence\n3. Record the decision'),20)
        self.assertTrue(any(c['question']=='What does HTTP stand for?' for c in cards))
        sequence=next(c for c in cards if c['kind']=='process')
        self.assertTrue(sequence['question'].endswith('in order?'))
        self.assertEqual(sequence['answer'].splitlines()[1],'2. Check the evidence')

    def test_powerpoint_table_retains_column_context(self):
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[5]);slide.shapes.title.text='Algorithms'
        table=slide.shapes.add_table(2,3, Inches(1), Inches(2), Inches(7), Inches(2)).table
        for row,cells in zip(table.rows,[('Method','Time complexity','Space complexity'),('Binary search','O(log n)','O(1)')]):
            for cell,value in zip(row.cells,cells):cell.text=value
        stream=io.BytesIO();ppt.save(stream)
        cards=extractive(extract_document(stream.getvalue(),'table.pptx'),20)
        self.assertEqual(len(cards),2)
        self.assertTrue(all('Binary search' in c['question'] and 'Slide 1 table' in c['source']['label'] for c in cards))

    def test_no_unsupported_reasoning_from_requested_difficulty(self):
        cards=extractive(document('A list is an ordered mutable collection of items.'),20,'advanced')
        self.assertEqual(len(cards),1)
        self.assertNotEqual(cards[0]['kind'],'cause')

    def test_focus_selects_relevant_topic_without_external_search(self):
        cards=extractive(document('A variable is a name that refers to an object.',
                                  'A mitochondrion is an organelle that produces energy.',
                                  'An enzyme is a protein that catalyzes reactions.'),1,focus='mitochondrion energy')
        self.assertIn('mitochondrion',cards[0]['question'])

    def test_large_input_is_bounded_and_covers_later_sections(self):
        doc=document(*(f'Concept{i} is a distinct technical definition for unit{i}.' for i in range(1800)))
        start=time.monotonic();cards=extractive(doc,20)
        self.assertLess(time.monotonic()-start,8)
        self.assertEqual(len(cards),20)
        self.assertTrue(any(int(c['source']['label'].split()[-1])>1200 for c in cards))


class CloudTests(unittest.TestCase):
    def setUp(self):
        self.doc=document('A list is an ordered mutable collection that can contain duplicate items.')
        self.raw={'question':'What is a list?', 'answer':'An ordered mutable collection that can contain duplicate items.',
                  'quote':self.doc.units[0].text, 'source_id':'S1', 'kind':'definition'}

    def test_two_pass_generation_and_no_secret_in_prompts(self):
        with patch('backend.cloud._request', return_value=[self.raw]) as request:
            cards,warnings=cloud.generate(self.doc,10,'mixed','private-test-key')
        self.assertEqual(len(cards),1);self.assertEqual(request.call_count,2)
        self.assertFalse(warnings)
        for call in request.call_args_list:
            payload,key,timeout=call.args
            self.assertIs(payload['store'],False)
            self.assertTrue(payload['text']['format']['strict'])
            self.assertNotIn('private-test-key',json.dumps(payload))
            self.assertNotIn('tools',payload)
        self.assertIn('Review every draft',request.call_args.args[0]['input'][0]['content'])

    def test_bad_quotes_unresolved_prompts_duplicates_and_styles(self):
        sources,_=cloud.source_passages(self.doc)
        bad=[{**self.raw,'quote':'An invented statement about something else.'},
             {**self.raw,'source_id':'made-up'},
             {**self.raw,'question':'What do the notes say about lists?'},
             {**self.raw,'question':'What is it?'},
             {**self.raw,'answer':'x '*81}]
        self.assertEqual(cloud.verified(bad,sources,20,'mixed'),[])
        self.assertEqual(len(cloud.verified([self.raw,self.raw],sources,20,'mixed')),1)
        self.assertEqual(cloud.verified([self.raw],sources,20,'cloze'),[])

    def test_reviewer_failure_never_silently_returns_unreviewed_cards(self):
        with patch('backend.cloud._request',side_effect=[[self.raw],UserError('Timeout',504,'cloud_timeout')]):
            with self.assertRaises(UserError) as caught:cloud.generate(self.doc,10,'mixed','test')
        self.assertEqual(caught.exception.code,'cloud_timeout')

    def test_missing_key_never_calls_network(self):
        with patch('backend.cloud._request') as request:
            with self.assertRaises(UserError):cloud.generate(self.doc,10,'mixed','')
        request.assert_not_called()

    def test_long_document_selection_is_bounded_and_disclosed(self):
        doc=document(*(f'Chapter {i}\n'+('A list is a mutable collection of elements. '*80) for i in range(90)))
        selected,warnings=cloud.source_passages(doc)
        self.assertLessEqual(sum(len(s['text']) for s in selected),cloud.MAX_SOURCE_CHARS)
        self.assertTrue(warnings)
        self.assertTrue(any(any(int(n)>75 for n in __import__('re').findall(r'Paragraph (\d+)',s['label'])) for s in selected))

    def test_wire_contract_timeout_refusal_incomplete_and_http_errors(self):
        class Response:
            def __init__(self,data):self.data=json.dumps(data).encode()
            def __enter__(self):return self
            def __exit__(self,*args):pass
            def read(self,limit):return self.data
        success={'status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':json.dumps({'cards':[self.raw]})}]}]}
        with patch('backend.cloud.build_opener') as opener:
            opener.return_value.open.return_value=Response(success)
            self.assertEqual(len(cloud._request({'model':'gpt-4.1-mini'},'test-only',1)),1)
            request=opener.return_value.open.call_args.args[0]
            self.assertEqual(request.full_url,'https://api.openai.com/v1/responses')
            self.assertEqual(request.get_header('Authorization'),'Bearer test-only')
            for data,code in [({'status':'incomplete'},'incomplete_generation'),
                              ({'status':'completed','output':[{'type':'message','content':[{'type':'refusal'}]}]},'generation_refused')]:
                opener.return_value.open.return_value=Response(data)
                with self.assertRaises(UserError) as caught:cloud._request({},'test',1)
                self.assertEqual(caught.exception.code,code)
            for status,code in [(401,'cloud_auth'),(429,'cloud_limit'),(500,'cloud_unavailable')]:
                opener.return_value.open.side_effect=HTTPError(cloud.ENDPOINT,status,'Secret response text',{},None)
                with self.assertRaises(UserError) as caught:cloud._request({},'test',1)
                self.assertEqual(caught.exception.code,code)
                self.assertNotIn('Secret',str(caught.exception))


class ConsentTests(unittest.TestCase):
    def test_cloud_requires_explicit_consent_and_server_key(self):
        with tempfile.TemporaryDirectory() as tmp:
            app=create_app({'TESTING':True,'DATA_DIR':tmp,'SECRET_KEY':'test','OPENAI_API_KEY':'private-test-key','EXTRACTOR':lambda *_: document('A list is a mutable collection of values.')})
            client=app.test_client();config=client.get('/api/config')
            self.assertNotIn('private-test-key',config.get_data(as_text=True))
            self.assertTrue(config.json['cloud']['configured'])
            def post(consent=None):
                fields={'file':(io.BytesIO(b'fixture'),'x.docx'),'name':'Test','category':'Study','engine':'cloud'}
                if consent:fields['cloud_consent']=consent
                return client.post('/api/generate',data=fields,headers={'X-Aura-CSRF':config.json['csrf_token']})
            with patch('app.cloud_generate') as generator:
                self.assertEqual(post().json['code'],'cloud_consent_required')
                generator.assert_not_called()
                app.config['OPENAI_API_KEY']=''
                self.assertEqual(post('yes').json['code'],'cloud_not_configured')
                generator.assert_not_called()
                app.config['OPENAI_API_KEY']='private-test-key'
                generator.return_value=([{'question':'What is a list?','answer':'A mutable collection.','kind':'definition','source':{'label':'P1','excerpt':'A list is a mutable collection of values.'}}],[])
                result=post('yes');self.assertEqual(result.status_code,200,result.json)
                self.assertEqual(result.json['card_types'],{'definition':1})

    def test_offline_does_not_use_cloud_even_when_key_is_configured(self):
        with tempfile.TemporaryDirectory() as tmp:
            app=create_app({'TESTING':True,'DATA_DIR':tmp,'SECRET_KEY':'test','OPENAI_API_KEY':'test','EXTRACTOR':lambda *_: document('A list is a mutable collection of values.')})
            client=app.test_client();token=client.get('/api/config').json['csrf_token']
            with patch('app.cloud_generate') as generator,patch('backend.cloud.build_opener') as network:
                r=client.post('/api/generate',data={'file':(io.BytesIO(b'x'),'x.docx'),'name':'Deck','category':'Python'},headers={'X-Aura-CSRF':token})
                self.assertEqual(r.status_code,200);generator.assert_not_called();network.assert_not_called()


if __name__=='__main__':unittest.main()
