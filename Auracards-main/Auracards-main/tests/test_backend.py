import io
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest
from contextlib import closing
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from zipfile import ZipFile, ZIP_DEFLATED
from pypdf import PdfWriter
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from app import create_app
from backend.errors import UserError
from backend.extraction import extract_document, extract_isolated, MAX_FILE_BYTES, DocumentText, TextUnit
from backend.generation import extractive, ollama, _local_base
from fixtures import FACTS, pdf_bytes, docx_bytes, pptx_bytes


class ExtractionTests(unittest.TestCase):
    def test_three_formats_with_isolated_parser_and_grounded_cards(self):
        for kind, factory in [('pdf', pdf_bytes), ('docx', docx_bytes), ('pptx', pptx_bytes)]:
            with self.subTest(format=kind):
                document=extract_isolated(factory(), 'notes.'+kind)
                cards=extractive(document, 10, 'mixed')
                self.assertEqual(len(cards), 10)
                for card in cards:
                    self.assertIn(card['answer'].casefold().rstrip('.'), card['source']['excerpt'].casefold())
                    self.assertIn(card['source']['excerpt'], FACTS)
                    self.assertNotIn('notes say', card['question'])
                    self.assertTrue(any(card['source']['excerpt'] in u.text.replace('\n',' ') for u in document.units))
                    self.assertTrue(card['source']['label'].startswith({'pdf':'Page','docx':'Paragraph','pptx':'Slide'}[kind]))

    def test_invalid_uploads(self):
        for raw,name,code in [(b'', 'x.pdf','empty_file'), (b'hello','x.txt','unsupported_type'), (b'not pdf','x.pdf','wrong_format'), (b'not zip','x.docx','wrong_format'), (b'x'*(MAX_FILE_BYTES+1),'x.pdf','file_too_large')]:
            with self.subTest(name=name,code=code), self.assertRaises(UserError) as caught:
                extract_document(raw,name)
            self.assertEqual(caught.exception.code,code)

    def test_image_only_and_encrypted_pdf(self):
        w=PdfWriter();w.add_blank_page(width=100,height=100);stream=io.BytesIO();w.write(stream)
        with self.assertRaises(UserError) as caught:extract_document(stream.getvalue(),'scan.pdf')
        self.assertEqual(caught.exception.code,'no_text')
        w.encrypt('password');stream=io.BytesIO();w.write(stream)
        with self.assertRaisesRegex(UserError,'Password-protected'):extract_document(stream.getvalue(),'locked.pdf')

    def test_archive_guards(self):
        for entry,value in [('word/vbaProject.bin',b'x'), ('word/document.xml',b'<!DOCTYPE a>'), ('word/large.bin',b'x'*(2*1024*1024))]:
            stream=io.BytesIO()
            with ZipFile(stream,'w',ZIP_DEFLATED) as z:
                z.writestr('[Content_Types].xml','<Types/>')
                if entry!='word/document.xml':z.writestr('word/document.xml','<document/>')
                z.writestr(entry,value)
            with self.subTest(entry=entry),self.assertRaises(UserError):extract_document(stream.getvalue(),'guard.docx')

    def test_tables_and_slide_notes(self):
        doc=Document();table=doc.add_table(rows=1, cols=2);table.cell(0,0).text='Python types';table.cell(0,1).text=FACTS[0]
        stream=io.BytesIO();doc.save(stream)
        self.assertIn('Table 1 row 1',extract_document(stream.getvalue(),'table.docx').units[0].label)
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[5]);slide.shapes.title.text='Python';slide.notes_slide.notes_text_frame.text=FACTS[1]
        stream=io.BytesIO();ppt.save(stream)
        self.assertTrue(any(u.label=='Slide 1 notes' and FACTS[1] in u.text for u in extract_document(stream.getvalue(),'notes.pptx').units))

    def test_worker_error_and_timeout_cleanup(self):
        with self.assertRaises(UserError) as caught:extract_isolated(b'not pdf','bad.pdf')
        self.assertEqual(caught.exception.code,'wrong_format')
        with self.assertRaises(UserError) as caught:extract_isolated(pdf_bytes(),'notes.pdf',timeout=0)
        self.assertEqual(caught.exception.code,'extraction_timeout')
        import multiprocessing
        self.assertFalse(multiprocessing.active_children())

    def test_fewer_facts_deduplicate_and_difficulty(self):
        doc=DocumentText([TextUnit('\n'.join([FACTS[0],FACTS[0],FACTS[1]]),'Page 1')],[])
        cards=extractive(doc,50,'beginner',style='cloze');self.assertEqual(len(cards),2)
        self.assertTrue(all('_____' in c['question'] for c in cards))
        with self.assertRaises(UserError):extractive(DocumentText([TextUnit('Python\na = 1','Page 1')],[]))


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.now=10000.0
        self.app=create_app({'TESTING':True,'DATA_DIR':self.tmp.name,'SECRET_KEY':'test-only-secret','EXTRACTOR':extract_document,'FOCUS_CLOCK':lambda:self.now})
        self.client,self.csrf=self.new_client()
    def tearDown(self):self.tmp.cleanup()
    def new_client(self):
        client=self.app.test_client();config=client.get('/api/config')
        self.assertEqual(config.status_code,200)
        return client,config.json['csrf_token']
    def upload(self,raw=None,filename='notes.pdf',client=None,csrf=None,**fields):
        data={'file':(io.BytesIO(raw if raw is not None else pdf_bytes()),filename),'name':'Python Notes','category':'Python Basics','count':'10','difficulty':'mixed','engine':'extractive',**fields}
        return (client or self.client).post('/api/generate',data=data,headers={'X-Aura-CSRF':csrf or self.csrf},content_type='multipart/form-data')
    def create_focus(self,**kwargs):
        response=self.client.post('/api/focus-sessions',json={'deck_id':'private-deck-id','duration_minutes':1,'timer_url':'http://localhost/pomodoro',**kwargs},headers={'X-Aura-CSRF':self.csrf})
        self.assertEqual(response.status_code,201,response.json);return response.json
    def control(self,s,action,**kwargs):
        return self.client.post('/api/focus-sessions/'+s['session_id'],json={'action':action},headers={'Authorization':'Bearer '+s['completion_token'],**kwargs})
    def test_static_and_api(self):
        for route in ['/','/app.js','/styles.css','/pomodoro','/pomodoro.js','/pomodoro.css','/api/health']:
            response=self.client.get(route);self.assertEqual(response.status_code,200,route);response.close()
        self.assertEqual(self.client.get('/not-here').status_code,404)
    def test_generation_three_formats(self):
        for kind,factory in [('pdf',pdf_bytes),('docx',docx_bytes),('pptx',pptx_bytes)]:
            response=self.upload(factory(),'notes.'+kind)
            self.assertEqual(response.status_code,200,response.json);self.assertEqual(len(response.json['cards']),10)
        self.assertEqual(sorted(p.name for p in Path(self.tmp.name).iterdir()),['focus.sqlite3','study-library.sqlite3'])
        with closing(sqlite3.connect(Path(self.tmp.name)/'focus.sqlite3')) as db:
            self.assertEqual(db.execute('SELECT count(*) FROM focus_sessions').fetchone()[0],0)
    def test_generation_validation_and_error_messages(self):
        for fields,code in [({'count':'51'},'invalid_count'),({'count':'no'},'invalid_count'),({'name':''},'invalid_fields'),({'difficulty':'wrong'},'invalid_options'),({'engine':'remote-ai'},'invalid_options')]:
            client,csrf=self.new_client();r=self.upload(client=client,csrf=csrf,**fields)
            self.assertEqual(r.status_code,400);self.assertEqual(r.json['code'],code)
        r=self.client.post('/api/generate',headers={'X-Aura-CSRF':self.csrf});self.assertEqual(r.json['code'],'missing_file')
        client,csrf=self.new_client();r=self.upload(docx_bytes([FACTS[0]]),'short.docx',client=client,csrf=csrf,count='50')
        self.assertEqual(r.status_code,200);self.assertEqual(r.json['generated_count'],1);self.assertTrue(r.json['warnings'])
    def test_csrf_origin_and_host(self):
        self.assertEqual(self.client.post('/api/generate').status_code,403)
        r=self.client.post('/api/focus-sessions',json={},headers={'X-Aura-CSRF':self.csrf,'Origin':'https://evil.example'})
        self.assertEqual(r.json['code'],'origin_blocked')
        self.assertEqual(self.client.get('/api/config',headers={'Origin':'https://evil.example'}).status_code,403)
        self.assertEqual(self.client.get('/api/health',headers={'Host':'evil.example'}).status_code,400)
        cookie=self.client.get('/api/config').headers.get('Set-Cookie')
        # Cookie is issued when the device is first initialized.
        c=self.app.test_client().get('/api/config').headers['Set-Cookie'];self.assertIn('HttpOnly',c);self.assertIn('SameSite=Strict',c);self.assertIn('Expires=',c)
        response=self.client.get('/')
        self.assertIn("script-src 'self'",response.headers['Content-Security-Policy']);response.close()
    def test_rate_limit(self):
        for _ in range(6):self.assertEqual(self.upload(raw=b'bad').status_code,415)
        self.assertEqual(self.upload().json['code'],'rate_limited')

    def test_busy_limit_and_release(self):
        entered,released=threading.Event(),threading.Event()
        def blocked(raw,name):
            entered.set()
            if not released.wait(5):raise RuntimeError('Test worker did not receive release.')
            return extract_document(raw,name)
        self.app.config['EXTRACTOR']=blocked
        worker,worker_csrf=self.new_client();results=[]
        thread=threading.Thread(target=lambda:results.append(self.upload(client=worker,csrf=worker_csrf)))
        thread.start()
        try:
            self.assertTrue(entered.wait(3))
            self.assertEqual(self.upload().json['code'],'server_busy')
        finally:released.set();thread.join(5)
        self.assertEqual(results[0].status_code,200)
        self.app.config['EXTRACTOR']=extract_document
        self.assertEqual(self.upload().status_code,200)
    def test_focus_lifecycle_pause_no_spoof_and_idempotency(self):
        s=self.create_focus();path='/api/focus-sessions/'+s['session_id'];header={'Authorization':'Bearer '+s['completion_token']}
        self.assertEqual(self.control(s,'complete').status_code,409)
        self.assertEqual(self.control(s,'start').json['status'],'running')
        self.now+=20;self.assertEqual(self.control(s,'pause').json['elapsed_seconds'],20)
        self.now+=1200;self.assertEqual(self.client.get(path,headers=header).json['remaining_seconds'],40)
        self.assertEqual(self.control(s,'complete').status_code,409)
        self.control(s,'start');self.now+=39;self.assertEqual(self.control(s,'complete').status_code,409)
        self.now+=1;self.assertEqual(self.control(s,'complete').json['status'],'completed')
        self.control(s,'complete');summary=self.client.get('/api/focus-sessions').json
        self.assertEqual(summary['completed_sessions'],1);self.assertEqual(summary['total_minutes'],1)
        other,_=self.new_client();self.assertEqual(other.get('/api/focus-sessions').json['total_minutes'],0)
        self.client.delete('/api/focus-sessions',headers={'X-Aura-CSRF':self.csrf});self.assertEqual(self.client.get('/api/focus-sessions').json['total_minutes'],0)
        self.assertEqual(self.client.get(path,headers=header).status_code,401)
    def test_focus_token_and_external_origin(self):
        s=self.create_focus(timer_url='https://timer.example/focus');path='/api/focus-sessions/'+s['session_id']
        self.assertEqual(self.client.get(path).status_code,401)
        self.assertEqual(self.client.get(path,headers={'Authorization':'Bearer wrong'}).status_code,401)
        r=self.control(s,'start',Origin='https://bad.example');self.assertEqual(r.status_code,403)
        r=self.control(s,'start',Origin='https://timer.example');self.assertEqual(r.status_code,200)
        self.assertEqual(r.headers['Access-Control-Allow-Origin'],'https://timer.example')
        pre=self.client.options(path,headers={'Origin':'https://timer.example','Access-Control-Request-Headers':'Authorization, Content-Type'})
        self.assertEqual(pre.status_code,204)
        self.assertEqual(self.control(s,'cancel').json['status'],'cancelled')
        self.assertEqual(self.control(s,'start').status_code,409)
        self.now+=48*3600+1;self.assertEqual(self.client.get(path,headers={'Authorization':'Bearer '+s['completion_token']}).status_code,410)
    def test_focus_input_and_open_limit(self):
        for fields in [{'duration_minutes':True},{'duration_minutes':121},{'deck_id':''},{'timer_url':'http://remote.example'},{'timer_url':'http://[broken'},{'timer_url':'http://localhost:bad'}]:
            r=self.client.post('/api/focus-sessions',json={'deck_id':'d','duration_minutes':1,**fields},headers={'X-Aura-CSRF':self.csrf})
            self.assertEqual(r.status_code,400,r.json)
        for _ in range(10):self.create_focus()
        r=self.client.post('/api/focus-sessions',json={'deck_id':'d','duration_minutes':1},headers={'X-Aura-CSRF':self.csrf})
        self.assertEqual(r.status_code,429)
    def test_no_uploaded_text_in_database_or_token_plaintext(self):
        self.upload();s=self.create_focus()
        raw=(Path(self.tmp.name)/'focus.sqlite3').read_bytes()
        self.assertNotIn(FACTS[0].encode(),raw);self.assertNotIn(s['completion_token'].encode(),raw)


class OllamaTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.mode='good';cls.payload=None
        class Handler(BaseHTTPRequestHandler):
            def log_message(self,*args):pass
            def do_GET(self):self.reply({'models':[{'name':'test:local'}]})
            def do_POST(self):
                body=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                if self.path=='/api/show':self.reply({'remote_host':'remote.example'} if cls.mode=='remote' else {})
                else:
                    cls.payload=body
                    if cls.mode=='redirect':
                        self.send_response(307);self.send_header('Location','https://remote.example/inference');self.end_headers();return
                    if cls.mode=='badjson':self.reply({'message':{'content':'not json'}});return
                    quote=FACTS[0] if cls.mode!='badquote' else 'This quote does not occur in the supplied source.'
                    cards=[{'question':'What is a variable in Python?','answer':FACTS[0],'quote':quote}]*2
                    self.reply({'message':{'content':json.dumps({'cards':cards})}})
            def reply(self,data):
                raw=json.dumps(data).encode();self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),Handler);cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
        cls.base='http://127.0.0.1:'+str(cls.server.server_port)
    @classmethod
    def tearDownClass(cls):cls.server.shutdown();cls.server.server_close();cls.thread.join()
    def setUp(self):type(self).mode='good';self.doc=DocumentText([TextUnit(FACTS[0],'Paragraph 1')],[])
    def test_structured_output_source_verification_and_deduplication(self):
        cards,warnings=ollama(self.doc,5,'mixed',self.base,'test:local');self.assertEqual(len(cards),1)
        self.assertEqual(cards[0]['source']['excerpt'],FACTS[0]);self.assertIsInstance(self.payload['format'],dict);self.assertFalse(self.payload['stream'])
        self.assertEqual(self.payload['options']['temperature'],0)
    def test_reject_unsupported_quotes_and_malformed_output(self):
        for mode,code in [('badquote','no_cards'),('badjson','invalid_generation'),('remote','local_model_required'),('redirect','redirect_blocked')]:
            type(self).mode=mode
            with self.subTest(mode=mode),self.assertRaises(UserError) as caught:ollama(self.doc,5,'mixed',self.base,'test:local')
            self.assertEqual(caught.exception.code,code)
    def test_loopback_only_and_cloud_model_rejected(self):
        for url in ['https://api.example','http://example.com','http://user:secret@localhost:11434','http://localhost:11434?x=y']:
            with self.assertRaises(ValueError):_local_base(url)
        with self.assertRaises(UserError) as caught:ollama(self.doc,5,'mixed',self.base,'test-cloud')
        self.assertEqual(caught.exception.code,'local_model_required')
        with self.assertRaises(UserError) as caught:ollama(self.doc,5,'mixed',self.base,'missing:local')
        self.assertEqual(caught.exception.code,'model_missing')

if __name__=='__main__':unittest.main()
