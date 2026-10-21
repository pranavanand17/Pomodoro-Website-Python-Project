"""End-to-end regressions for real study-document layouts and answer pairing."""
import io
import tempfile
import unittest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app import create_app
from backend.errors import UserError
from backend.extraction import DocumentText, TextUnit, extract_document, normalize, join_pdf_lines
from backend.generation import extractive, extractive_details
from fixtures import pdf_bytes, docx_bytes


def cards(text, **options):
    return extractive(DocumentText([TextUnit(text,'Page 1')],[]),50,**options)


class StructuredTests(unittest.TestCase):
    def test_label_variants(self):
        for q,a in [('Q:','A:'),('Q1:','Ans:'),('Q1.','Ans.'),('Question 1:','Answer:'),('1. Question:','Answer 1:'),('Q2)','A)')]:
            with self.subTest(q=q):
                result=cards(f'{q} What does len return?\n{a} The number of items.')
                self.assertEqual(len(result),1)
                self.assertEqual(result[0]['question'],'What does len return?')
                self.assertEqual(result[0]['answer'],'The number of items.')
                self.assertEqual(result[0]['origin'],'qa')

    def test_unlabelled_and_numbered_questions_preserve_answers(self):
        result=cards('1. What is a list?\nA mutable ordered collection.\nIt permits duplicates.\n2. What is a tuple?\nAn immutable ordered collection.')
        self.assertEqual(len(result),2)
        self.assertEqual(result[0]['answer'],'A mutable ordered collection.\nIt permits duplicates.')
        self.assertEqual(result[1]['question'],'What is a tuple?')
        self.assertNotIn('immutable',result[0]['answer'])

    def test_inline_and_wrapped_question(self):
        result=cards('Q: What is Python? A: A programming language.\nQ: How does a function\nreturn a value?\nAnswer: With the return statement.')
        self.assertEqual(len(result),2)
        self.assertEqual(result[1]['question'],'How does a function return a value?')
        self.assertEqual(result[0]['answer'],'A programming language.')

    def test_label_only_and_imperative_question(self):
        result=cards('Question:\nWhat is a tuple?\nAnswer:\nAn immutable sequence.\nQ: Define encapsulation.\nA: Bundling state and behavior together.')
        self.assertEqual(len(result),2)
        self.assertEqual(result[0]['question'],'What is a tuple?')
        self.assertEqual(result[1]['question'],'Define encapsulation.')

    def test_multiline_lists_and_code_are_not_cut_or_rewritten(self):
        result=cards('Q: How can a function return the square of x?\nA:\ndef square(x):\n    return x * x\nQ: Which collection types are listed?\nA:\n- list\n- tuple\n- set')
        self.assertEqual(len(result),2)
        self.assertEqual(result[0]['answer'],'def square(x):\n    return x * x')
        self.assertEqual(result[1]['answer'],'- list\n- tuple\n- set')
        self.assertIn('    return',normalize('def square(x):\n    return x * x'))

    def test_questions_without_answers_do_not_borrow_the_next_answer(self):
        doc=DocumentText([TextUnit('Q: What is a generator?\nQ: What is a tuple?\nA: An immutable sequence.','P1')],[])
        result,notes,meta=extractive_details(doc,20)
        self.assertEqual(len(result),1)
        self.assertEqual(result[0]['question'],'What is a tuple?')
        self.assertTrue(any('no detected answer' in w for w in notes))

    def test_short_numeric_and_literal_answers(self):
        result=cards('Q: What value means no result?\nA: None\nQ: What is the first index of a Python list?\nA: 0\nQ: Can lists contain duplicates?\nA: Yes')
        self.assertEqual([c['answer'] for c in result],['None','0','Yes'])

    def test_topics_with_explicit_explanations(self):
        result=cards('Topic: Inheritance\nExplanation: A class reuses behavior from another class.\nIt may override inherited methods.\nTopic: Python lists\nExplanation: Lists are mutable ordered collections.')
        self.assertEqual(len(result),2)
        self.assertEqual(result[0]['question'],'What is Inheritance?')
        self.assertEqual(result[1]['question'],'What are Python lists?')
        self.assertIn('It may override',result[0]['answer'])
        self.assertNotIn('Lists',result[0]['answer'])

    def test_plain_markdown_and_colon_headings(self):
        for title in ['Inheritance','## Inheritance','Inheritance:','**Inheritance**']:
            with self.subTest(title=title):
                result=cards(title+'\nA class reuses behavior from another class.\nPolymorphism\nDifferent objects respond to the same interface.')
                self.assertEqual(len(result),2)
                self.assertTrue(all(c['origin']=='topic' for c in result))

    def test_plural_steps_and_comparison_titles(self):
        result=cards('Advantages of functions\nThey support reuse and organize programs.\nLists vs tuples\nLists are mutable; tuples are immutable.\nSteps of the review process:\n1. Read the request\n2. Check the evidence\n3. Record the decision',input_format='topics')
        self.assertEqual(len(result),3)
        self.assertEqual(result[0]['question'],'What are the advantages of functions?')
        self.assertEqual(result[1]['question'],'How do Lists and tuples differ?')
        self.assertTrue(result[2]['question'].endswith('in order?'))

    def test_explicit_modes_do_not_generate_other_types(self):
        text='Q: What is a tuple?\nA: An immutable sequence.\nTopic: Inheritance\nA class reuses behavior from a base class.'
        self.assertTrue(all(c['origin']=='qa' for c in cards(text,input_format='qa')))
        self.assertTrue(all(c['origin']=='topic' for c in cards(text,input_format='topics')))
        with self.assertRaises(UserError):cards('A list is a mutable sequence.',input_format='qa')

    def test_topics_mode_accepts_lower_case_heading(self):
        result=cards('exception handling\nA program can recover from certain errors using try and except.',input_format='topics')
        self.assertEqual(result[0]['question'],'What is exception handling?')

    def test_structured_cards_are_prioritized_and_answers_not_reused(self):
        text='Q: What does len return?\nA: The number of items.\nTopic: Inheritance\nA class reuses behavior from a base class.'
        result=cards(text)
        self.assertEqual(len(result),2)
        self.assertEqual({c['origin'] for c in result},{'qa','topic'})
        self.assertEqual(len({c['question'] for c in result}),2)

    def test_long_blocks_are_skipped_intact_with_diagnostic(self):
        doc=DocumentText([TextUnit('Q: Explain the process?\nA: '+('This sentence is part of the answer. '*90)+'\nQ: What is a tuple?\nA: An immutable sequence.','P1')],[])
        result,notes,_=extractive_details(doc,20)
        self.assertEqual(len(result),1)
        self.assertTrue(any('length limit' in n for n in notes))

    def test_cloze_preference_does_not_destroy_authored_pairs(self):
        doc=DocumentText([TextUnit('Q: What is a tuple?\nA: An immutable sequence.','P1')],[])
        result,notes,meta=extractive_details(doc,20,style='cloze')
        self.assertEqual(result[0]['question'],'What is a tuple?')
        self.assertEqual(meta,{'detected_format':'qa','structured_cards':1})
        self.assertTrue(notes)

    def test_pdf_preserves_long_question_answer_boundary(self):
        text='Q: Explain the role of the Python interpreter in executing source code\nAns: It executes the program.'
        self.assertIn('\nAns:',join_pdf_lines(text))

    def test_same_pair_in_real_pdf_docx_pptx(self):
        lines=['Q: What is a tuple?','A: An immutable sequence.','Q: What does len return?','A: The number of items.']
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[5]);slide.shapes.title.text='Python review'
        slide.shapes.add_textbox(Inches(.5),Inches(1),Inches(8),Inches(5)).text_frame.text='\n'.join(lines)
        data=io.BytesIO();ppt.save(data)
        for kind,raw in [('pdf',pdf_bytes(lines)),('docx',docx_bytes(lines)),('pptx',data.getvalue())]:
            with self.subTest(kind=kind):
                result=extractive(extract_document(raw,'qa.'+kind),20)
                self.assertEqual([(c['question'],c['answer']) for c in result],[('What is a tuple?','An immutable sequence.'),('What does len return?','The number of items.')])

    def test_word_bold_headings_and_paragraph_answers(self):
        doc=Document();doc.add_paragraph().add_run('Inheritance').bold=True
        doc.add_paragraph('A class reuses behavior from a base class.')
        doc.add_paragraph('It may override inherited methods.')
        doc.add_heading('Python lists',1);doc.add_paragraph('Lists are mutable ordered collections.')
        out=io.BytesIO();doc.save(out);result=extractive(extract_document(out.getvalue(),'topics.docx'),20)
        self.assertEqual(len(result),2)
        self.assertIn('override',result[0]['answer'])

    def test_word_qa_and_topic_tables(self):
        for headers,values,origin in [(('Question','Answer'),('What is a tuple?','An immutable sequence.'),'qa'),(('Topic','Explanation'),('Inheritance','A class reuses behavior from another class.'),'topic')]:
            with self.subTest(origin=origin):
                doc=Document();table=doc.add_table(rows=2,cols=2)
                for row,items in zip(table.rows,[headers,values]):
                    for cell,value in zip(row.cells,items):cell.text=value
                out=io.BytesIO();doc.save(out);result=extractive(extract_document(out.getvalue(),'table.docx'),20)
                self.assertEqual(len(result),1);self.assertEqual(result[0]['origin'],origin)
                self.assertEqual(result[0]['answer'],values[1])

    def test_api_format_metadata_and_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            app=create_app({'TESTING':True,'DATA_DIR':tmp,'SECRET_KEY':'test','EXTRACTOR':extract_document})
            client=app.test_client();csrf=client.get('/api/config').json['csrf_token']
            def post(fmt):
                return client.post('/api/generate',data={'file':(io.BytesIO(docx_bytes(['Q: What is a tuple?','A: An immutable sequence.'])),'notes.docx'),'name':'Deck','category':'Python','input_format':fmt},headers={'X-Aura-CSRF':csrf})
            good=post('qa');self.assertEqual(good.status_code,200,good.json)
            self.assertEqual(good.json['detected_format'],'qa');self.assertEqual(good.json['structured_cards'],1)
            self.assertEqual(post('invalid').status_code,400)


if __name__=='__main__':unittest.main()
