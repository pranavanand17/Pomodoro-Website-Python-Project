"""Bounded, text-only extraction. No filenames or document contents are saved."""
from collections import Counter
from dataclasses import dataclass
from io import BytesIO
import multiprocessing
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from .errors import UserError

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNITS = 300
MAX_TEXT_CHARS = 300_000
MAX_ARCHIVE_BYTES = 80 * 1024 * 1024


@dataclass
class TextUnit:
    text: str
    label: str
    kind: str = "text"
    heading: str = ""
    columns: tuple = ()


@dataclass
class DocumentText:
    units: list[TextUnit]
    warnings: list[str]


def normalize(text):
    text = unicodedata.normalize("NFKC", text or "")
    text = re.sub(r"([A-Za-z])-\s*\n\s*([a-z])", r"\1\2", text)
    text = "".join(c for c in text if c in "\n\t" or unicodedata.category(c)[0] != "C")
    # Keep indentation in explicit Q&A code answers; normalize interior prose
    # spacing without flattening a Python code block into invalid code.
    lines=[]
    for line in text.expandtabs(4).splitlines():
        indent=len(line)-len(line.lstrip(' '))
        lines.append(' '*min(indent,40)+re.sub(r'[ \t]+',' ',line.lstrip(' ')).rstrip())
    return '\n'.join(lines).strip()


def join_pdf_lines(text):
    """Join visual wraps in prose, retaining short headings and list boundaries."""
    lines = []
    for line in text.splitlines():
        if (lines and len(lines[-1]) >= 45 and len(lines[-1].split()) >= 5
                and not re.search(r"[.!?:]$", lines[-1]) and line
                and not re.match(r"^(?:[•▪●*]|\d+[.)])\s+", line)
                and not re.match(r'^(?:Q(?:uestion)?\s*\d*|A(?:ns(?:wer)?)?\s*\d*|Topic|Explanation)\s*[.:–-]', line, re.I)
                and not (len(line) < 45 and len(line.split()) <= 7 and not re.search(r"[.!?=:]", line))):
            lines[-1] += " " + line
        else:
            lines.append(line)
    return "\n".join(lines)


def _archive_check(raw, kind):
    try:
        with ZipFile(BytesIO(raw)) as archive:
            infos = archive.infolist()
            if len(infos) > 2500 or sum(i.file_size for i in infos) > MAX_ARCHIVE_BYTES:
                raise UserError("The document expands beyond the processing limit. Split it into smaller files.")
            names = {i.filename for i in infos}
            required = "word/document.xml" if kind == "docx" else "ppt/presentation.xml"
            if required not in names or "[Content_Types].xml" not in names:
                raise UserError("The file contents do not match its document type.", 415, "wrong_format")
            for item in infos:
                if item.flag_bits & 1:
                    raise UserError("Password-protected documents are not supported. Upload an unlocked copy.")
                if "vbaproject" in item.filename.lower():
                    raise UserError("Macro-enabled documents are not supported.", 415, "macros")
                if item.file_size > 25 * 1024 * 1024 or item.file_size / max(item.compress_size, 1) > 1000:
                    raise UserError("The document contains an unusually large compressed entry.")
                if item.filename.endswith((".xml", ".rels")):
                    xml = archive.read(item)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise UserError("The document contains unsupported XML declarations.")
    except BadZipFile as exc:
        raise UserError("This is not a readable DOCX or PPTX file.", 415, "wrong_format") from exc


def _pdf(raw):
    from pypdf import PdfReader
    if not raw[:1024].lstrip().startswith(b"%PDF-"):
        raise UserError("The file is not a PDF document.", 415, "wrong_format")
    reader = PdfReader(BytesIO(raw), strict=False)
    if reader.is_encrypted:
        raise UserError("Password-protected PDFs are not supported. Upload an unlocked copy.")
    if len(reader.pages) > MAX_UNITS:
        raise UserError("PDFs are limited to 300 pages. Split this document first.")
    units, empty = [], 0
    for i, page in enumerate(reader.pages, 1):
        text = normalize(page.extract_text() or "")
        if text:
            units.append(TextUnit(text, f"Page {i}"))
        else:
            empty += 1
    # Remove repeated page headers/footers, without altering the middle of pages.
    if len(units) >= 3:
        edges = Counter()
        for unit in units:
            lines = unit.text.splitlines()
            edges.update(set(lines[:1] + lines[-1:]))
        repeated = {line for line, n in edges.items() if n >= max(3, len(units) * .7) and len(line) < 120}
        for unit in units:
            lines = unit.text.splitlines()
            unit.text = "\n".join(line for j, line in enumerate(lines)
                                  if not (line in repeated and j in (0, len(lines) - 1)))
    for unit in units:
        unit.text = join_pdf_lines(unit.text)
    warnings = [f"{empty} page(s) had no extractable text and were skipped."] if empty else []
    return DocumentText(units, warnings)


def _docx(raw):
    from docx import Document
    from docx.table import Table
    _archive_check(raw, "docx")
    doc = Document(BytesIO(raw))
    units, paragraph, table, heading = [], 0, 0, ""
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            table += 1
            headers = ()
            for i, row in enumerate(block.rows, 1):
                cells = [normalize(c.text).replace('\n', ' ') for c in row.cells]
                if i == 1 and _table_header(cells):
                    headers = tuple(cells)
                    continue
                text = " | ".join(cells)
                if text:
                    units.append(TextUnit(text, f"Table {table} row {i}", "table", heading, headers))
        else:
            paragraph += 1
            text = normalize(block.text)
            if text:
                nonempty=[r for r in block.runs if r.text.strip()]
                bold_heading = bool(nonempty and all(r.bold for r in nonempty) and len(text)<=100 and not re.search(r'[.!?]',text))
                is_heading = bold_heading or bool(block.style and (block.style.name.startswith('Heading') or block.style.name in ('Title', 'Subtitle')))
                if is_heading:
                    heading = text
                is_list = bool(block.style and block.style.name.startswith('List'))
                units.append(TextUnit(text, f"Paragraph {paragraph}", "heading" if is_heading else "list" if is_list else "text", heading))
    return DocumentText(units, [])


def _table_header(cells):
    known = {'term', 'concept', 'name', 'type', 'property', 'feature', 'definition', 'meaning',
             'description', 'function', 'value', 'unit', 'units', 'symbol', 'stage', 'step',
             'process', 'method', 'time complexity', 'space complexity', 'example',
             'question', 'answer', 'q', 'a', 'ans', 'prompt', 'response', 'topic', 'explanation'}
    return len(cells) >= 2 and all(len(c) < 80 for c in cells) and sum(c.casefold() in known for c in cells) >= 2


def _shape_content(shapes):
    for shape in shapes:
        if hasattr(shape, 'shapes'):
            yield from _shape_content(shape.shapes)
        if getattr(shape, 'has_text_frame', False):
            yield 'text', shape.text_frame.text, ()
        if getattr(shape, 'has_table', False):
            headers = ()
            for i, row in enumerate(shape.table.rows):
                cells = [normalize(c.text).replace('\n', ' ') for c in row.cells]
                if i == 0 and _table_header(cells):
                    headers = tuple(cells)
                    continue
                yield 'table', ' | '.join(cells), headers


def _pptx(raw):
    from pptx import Presentation
    _archive_check(raw, "pptx")
    presentation = Presentation(BytesIO(raw))
    if len(presentation.slides) > MAX_UNITS:
        raise UserError("Presentations are limited to 300 slides. Split this file first.")
    units = []
    for i, slide in enumerate(presentation.slides, 1):
        content = list(_shape_content(slide.shapes))
        text = normalize("\n".join(t for kind, t, _ in content if kind == 'text'))
        heading = normalize(slide.shapes.title.text) if slide.shapes.title is not None else ''
        if text:
            units.append(TextUnit(text, f"Slide {i}", heading=heading))
        row = 0
        for kind, value, columns in content:
            if kind == 'table' and value.strip():
                row += 1
                units.append(TextUnit(normalize(value), f'Slide {i} table row {row}', kind, heading, columns))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            note = normalize(frame.text if frame else "")
            if note:
                units.append(TextUnit(note, f"Slide {i} notes"))
    return DocumentText(units, [])


def extract_document(raw, filename):
    if not raw:
        raise UserError("The uploaded file is empty.", 400, "empty_file")
    if len(raw) > MAX_FILE_BYTES:
        raise UserError("Choose a document no larger than 20 MB.", 413, "file_too_large")
    kind = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if kind not in ("pdf", "docx", "pptx"):
        raise UserError("Supported formats are PDF, DOCX and PPTX.", 415, "unsupported_type")
    try:
        result = {"pdf": _pdf, "docx": _docx, "pptx": _pptx}[kind](raw)
    except UserError:
        raise
    except Exception as exc:
        raise UserError("The document could not be read. Check that it is valid and unlocked.") from exc
    result.units = [u for u in result.units if u.text.strip()]
    size = sum(len(u.text) for u in result.units)
    if size > MAX_TEXT_CHARS or len(result.units) > 6000:
        raise UserError("This document contains too much text. Split it into smaller sections.")
    if size < 40 or not any(len(re.findall(r"\w+", u.text)) >= 5 for u in result.units):
        raise UserError("No usable study text was found. Scanned documents need OCR; image-only slides are not supported.", 422, "no_text")
    return result


def _worker(conn, raw, filename):
    try:
        try:
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (1024 ** 3, 1024 ** 3))
            resource.setrlimit(resource.RLIMIT_CPU, (25, 25))
        except (ImportError, ValueError, OSError):
            pass
        conn.send((True, extract_document(raw, filename)))
    except UserError as exc:
        conn.send((False, (str(exc), exc.status, exc.code)))
    except Exception:
        conn.send((False, ("Document processing failed. Try a smaller, unlocked file.", 422, "processing_failed")))
    finally:
        conn.close()


def extract_isolated(raw, filename, timeout=35):
    """A separate process bounds parser time; Unix additionally bounds memory/CPU."""
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(target=_worker, args=(send, raw, filename), daemon=True)
    try:
        process.start()
        send.close()
        if not receive.poll(timeout):
            raise UserError("Document extraction timed out. Split the document and try again.", 422, "extraction_timeout")
        try:
            ok, value = receive.recv()
        except EOFError as exc:
            raise UserError("The document exceeded the parser's resource limit. Try a smaller file.") from exc
        if not ok:
            raise UserError(*value)
        return value
    finally:
        receive.close()
        send.close()
        if process.is_alive():
            process.terminate()
        if process.pid:
            process.join(timeout=3)
